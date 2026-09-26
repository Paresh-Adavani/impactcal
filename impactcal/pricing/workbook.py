import pandas as pd, csv, json, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
B='/home/claude/build/data/'
cost=pd.read_csv(B+'costing.csv',encoding='utf-8-sig').fillna('')
P=json.load(open('model_params.json'))['rubber']
wb=Workbook(); ws=wb.active; ws.title='Policy'
NAVY='1B3160'; hdr_fill=PatternFill('solid',fgColor=NAVY); hdr_font=Font(bold=True,color='FFFFFF'); edit_fill=PatternFill('solid',fgColor='FFF2CC'); est_fill=PatternFill('solid',fgColor='E8F0FE')
rows=[['ADONI TECH — ImpactCal global price list & costing','',''],['Prepared','2026-09-24',''],['',''],
['PRICING POLICY (agreed 23-24 Sep 2026)','',''],
['File prices','Highest price found in any of the 8 uploaded price lists is taken as-is (list price).',''],
['Estimated prices','Cost model per series (price ~ k · Fs^a · stroke^b fitted on your own prices; median error 6 % crane, 9 % industrial) + negotiation margin below',''],
['Rubber mounts','Moulding cost + hardware + mould cost amortised, + margin over cost (parameters below; Rubber mounts sheet recalculates live)',''],
['Rounding','nearest ₹50 below ₹10,000, nearest ₹100 above (always up)',''],
['USD','list ÷ live INR/USD rate × (1 + export uplift); recomputed at quotation time; column here is indicative',''],
['',''],
['PARAMETERS (yellow cells are yours to change)','Value','Unit'],
['Negotiation / dealer margin on estimated prices',25,'%'],
['Dealer price = list × (1 − discount)',20,'%'],
['INR per USD (indicative)',95.87,'INR'],
['Export uplift',12,'%'],
['Rubber compound',P['rubber_rate'],'₹/kg'],
['Rubber density',P['rubber_density'],'g/cm³'],
['Moulding cost per shot',P['moulding_shot'],'₹'],
['Moulding cost per kg rubber',P['moulding_per_kg'],'₹/kg'],
['Metal parts material',P['metal_rate'],'₹/kg'],
['Metal parts machining / plating',P['metal_per_kg_machining'],'₹/kg'],
['Minimum hardware per piece',P['hardware_min'],'₹'],
['Mould cost — small (Ø ≤ 80)',P['mould_small'],'₹'],
['Mould cost — medium (Ø ≤ 160)',P['mould_medium'],'₹'],
['Mould cost — large',P['mould_large'],'₹'],
['Cavities per mould — small / medium / large',f"{P['cav_small']} / {P['cav_medium']} / {P['cav_large']}",''],
['Mould amortised over',P['mould_qty'],'pcs'],
['Margin over cost (rubber mounts)',int(P['margin']*100),'%'],
]
for r in rows: ws.append(r)
ws['A1'].font=Font(bold=True,size=14,color=NAVY); ws.column_dimensions['A'].width=48; ws.column_dimensions['B'].width=90; ws.column_dimensions['C'].width=8
for i in range(12,29): ws.cell(i,2).fill=edit_fill
for i in (4,11): ws.cell(i,1).font=Font(bold=True,color=NAVY)
# named cells
NM={'margin_est':'Policy!$B$12','dealer_disc':'Policy!$B$13','usd_rate':'Policy!$B$14','uplift':'Policy!$B$15','rub_rate':'Policy!$B$16','rub_dens':'Policy!$B$17','shot':'Policy!$B$18','mould_kg':'Policy!$B$19','metal':'Policy!$B$20','mach':'Policy!$B$21','hw_min':'Policy!$B$22','m_s':'Policy!$B$23','m_m':'Policy!$B$24','m_l':'Policy!$B$25','qty':'Policy!$B$27','rmargin':'Policy!$B$28'}
def sheet(name, df, cols, widths):
    s=wb.create_sheet(name); s.append(cols)
    for c in s[1]: c.fill=hdr_fill; c.font=hdr_font; c.alignment=Alignment(wrap_text=True,vertical='center')
    for _,r in df.iterrows(): s.append([r.get(c,'') if c in df.columns else '' for c in cols])
    for i,w in enumerate(widths,1): s.column_dimensions[get_column_letter(i)].width=w
    s.freeze_panes='C2'; s.auto_filter.ref=f"A1:{get_column_letter(len(cols))}{len(df)+1}"
    return s
# --- shock absorbers / WRI / accessories / other: values, USD formula, dealer formula
def prod_sheet(name, table, extra_cols=()):
    d=cost[cost.table==table].copy()
    cols=['key','model','series','basis','file_price','file_source','estimate','margin_pct','list_inr','dealer_inr','usd_indicative','lead_time_days','note']
    s=sheet(name,d,cols,[14,26,14,8,11,26,11,9,11,11,11,8,60])
    n=len(d)
    for i in range(2,n+2):
        s.cell(i,11).value=f'=IF(I{i}>0,ROUNDUP(I{i}/{NM["usd_rate"]}*(1+{NM["uplift"]}/100),0),"")'
        s.cell(i,10).value=f'=IF(I{i}>0,CEILING(I{i}*(1-{NM["dealer_disc"]}/100),IF(I{i}<10000,50,100)),"")'
        s.cell(i,9).fill=edit_fill
        if s.cell(i,4).value!='file': s.cell(i,7).fill=est_fill
    return s
prod_sheet('Shock absorbers','shock_absorbers'); prod_sheet('Wire rope isolators','wire_rope_isolators'); prod_sheet('Accessories','accessories'); prod_sheet('Other products','other_products')
# --- rubber mounts with live cost formulas
rm=pd.DataFrame(list(csv.DictReader(open(B+'rubber_mounts.csv',encoding='utf-8-sig'))))
rmn=rm[['model','load_max_kg','height_mm','dia_mm','length_mm']].copy()
for c in ['load_max_kg','height_mm','dia_mm','length_mm']: rmn[c]=pd.to_numeric(rmn[c],errors='coerce')
d=cost[cost.table=='rubber_mounts'].merge(rmn,on='model',how='left')
for c in ['cavities','rubber_kg','metal_kg','load_max_kg']: d[c]=pd.to_numeric(d[c],errors='coerce')
d=d.where(pd.notna(d),'')
cols=['key','model','series','load_max_kg','dia_mm','height_mm','length_mm','cavities','rubber_kg','rubber_cost','moulding','metal_kg','hardware','mould_cost','mould_amort','unit_cost','margin_pct','list_inr','dealer_inr','usd_indicative','basis','note']
s=sheet('Rubber mounts',d,cols,[14,20,30,9,8,8,8,7,9,10,9,8,9,10,10,10,8,10,10,10,7,50])
for i in range(2,len(d)+2):
    # rubber cost = kg × rate; moulding = shot + kg × per-kg; hardware = max(min, metal_kg × (metal+machining)); mould amort = mould/(qty×cavities); unit = sum; list = ceiling(unit×(1+margin))
    s.cell(i,10).value=f'=ROUND(I{i}*{NM["rub_rate"]},0)'
    s.cell(i,11).value=f'=ROUND({NM["shot"]}+I{i}*{NM["mould_kg"]},0)'
    s.cell(i,13).value=f'=ROUND(MAX({NM["hw_min"]},L{i}*({NM["metal"]}+{NM["mach"]})),0)'
    s.cell(i,14).value=f'=IF(E{i}<=80,{NM["m_s"]},IF(E{i}<=160,{NM["m_m"]},{NM["m_l"]}))'
    s.cell(i,15).value=f'=ROUND(N{i}/({NM["qty"]}*MAX(1,H{i})),0)'
    s.cell(i,16).value=f'=J{i}+K{i}+M{i}+O{i}'
    if s.cell(i,21).value=='cost':
        s.cell(i,17).value=f'={NM["rmargin"]}'
        s.cell(i,18).value=f'=CEILING(P{i}*(1+Q{i}/100),IF(P{i}*(1+Q{i}/100)<10000,50,100))'
    else: s.cell(i,18).fill=edit_fill
    s.cell(i,19).value=f'=IF(R{i}>0,CEILING(R{i}*(1-{NM["dealer_disc"]}/100),IF(R{i}<10000,50,100)),"")'
    s.cell(i,20).value=f'=IF(R{i}>0,ROUNDUP(R{i}/{NM["usd_rate"]}*(1+{NM["uplift"]}/100),0),"")'
    for c in (8,9,12): s.cell(i,c).fill=edit_fill
# summary
sm=wb.create_sheet('Summary',1)
sm.append(['Table','Rows','From price lists','Estimated (model)','Costed (rubber)','Min list ₹','Max list ₹']); 
for c in sm[1]: c.fill=hdr_fill; c.font=hdr_font
for t,g in cost.groupby('table'):
    li=pd.to_numeric(g.list_inr,errors='coerce')
    sm.append([t,len(g),int((g.basis=='file').sum()),int((g.basis=='model').sum()),int((g.basis=='cost').sum()),int(li.min()) if li.notna().any() else '',int(li.max()) if li.notna().any() else ''])
sm.append([]); sm.append(['Legend','yellow = your input (list price / geometry)','blue = estimated value','',''])
for i,w in enumerate([22,8,16,16,14,12,12],1): sm.column_dimensions[get_column_letter(i)].width=w
out='/home/claude/build/pricing/ImpactCal_Pricelist_2026-09.xlsx'
import os; os.makedirs('/home/claude/build/pricing',exist_ok=True); wb.save(out); print(out, os.path.getsize(out))
